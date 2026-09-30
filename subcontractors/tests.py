from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import load_workbook

from suppliers.models import Supplier

from .models import DirectSubcontractor

EXPECTED_COLUMNS = [
	'Sl No', 'Vendor Code', 'Vendor / Firm Name', 'Trade', 'Service Type', 'Scope of Work',
	'Labour Strength Available (Nos)', 'Labour Licence No.', 'ESI Reg. No.', 'PF Reg. No.', 'Contact Person',
	'Mobile No.', 'Alternate No.', 'Email', 'Full Address', 'City', 'State', 'PIN Code', 'GST Registration Type',
	'GSTIN', 'GST State Code', 'PAN', 'GSTIN Check', 'MSME Category', 'Udyam Reg. No.', 'TDS Section',
	'Bank A/c Holder Name', 'Bank Name', 'Bank A/c No.', 'IFSC', 'Branch', 'IFSC Check', 'Referred By',
	'Worked with VISHWA Before?', 'Project(s) Worked On', 'Quality (1-5)', 'Timeliness (1-5)', 'Safety (1-5)',
	'Overall Rating', 'Last Reviewed On', 'Status', 'Remarks',
]


class SubcontractorUploadTests(TestCase):
	def setUp(self):
		self.admin = get_user_model().objects.create_superuser('boss', 'boss@example.com', 'pw')
		self.client.force_login(self.admin)

	def _template(self):
		response = self.client.get('/subcontractors/template.xlsx')
		self.assertEqual(response.status_code, 200)
		return response.content

	def test_template_has_the_workbook_columns_and_uploads(self):
		content = self._template()
		sheet = load_workbook(BytesIO(content))['Labour & Composite']
		rows = list(sheet.iter_rows(values_only=True))
		self.assertEqual(list(rows[0]), EXPECTED_COLUMNS)
		response = self.client.post('/subcontractors/upload/', {'file': SimpleUploadedFile('s.xlsx', content)})
		self.assertRedirects(response, '/subcontractors/')
		sub = DirectSubcontractor.objects.get(code='SC0001')
		self.assertEqual(sub.labour_strength, 40)
		self.assertEqual(sub.service_type, 'Labour')
		self.assertEqual(sub.timeliness_rating, 4)
		self.assertEqual(sub.safety_rating, 5)
		self.assertTrue(sub.gstin_check)
		self.assertEqual(sub.trade_links.get().trade.name, 'Civil')
		self.assertContains(self.client.get(f'/subcontractors/{sub.pk}/'), 'Shuttering and concreting labour')

	def test_upload_with_other_sheet_name_and_numeric_cells(self):
		workbook = load_workbook(BytesIO(self._template()))
		sheet = workbook['Labour & Composite']
		sheet.title = 'Sheet1'
		del workbook['Instructions']
		sheet['L2'] = 9876543210
		sheet['R2'] = 500072
		sheet['U2'] = 9
		sheet['AI2'] = 'Sreenidhi CR-2, A S Rao Nagar'
		output = BytesIO()
		workbook.save(output)
		self.client.post('/subcontractors/upload/', {'file': SimpleUploadedFile('s.xlsx', output.getvalue())})
		sub = DirectSubcontractor.objects.get(code='SC0001')
		self.assertEqual((sub.mobile_number, sub.pin_code, sub.gst_state_code), ('9876543210', '500072', '09'))
		self.assertEqual(sub.project_links.count(), 2)

	def test_invalid_rows_reject_the_whole_file(self):
		workbook = load_workbook(BytesIO(self._template()))
		sheet = workbook['Labour & Composite']
		bad_row = [''] * len(EXPECTED_COLUMNS)
		bad_row[:3] = [2, 'SC0002', 'Bad Ratings']
		bad_row[EXPECTED_COLUMNS.index('Labour Strength Available (Nos)')] = 'forty'
		bad_row[EXPECTED_COLUMNS.index('Quality (1-5)')] = 7
		sheet.append(bad_row)
		output = BytesIO()
		workbook.save(output)
		response = self.client.post('/subcontractors/upload/', {'file': SimpleUploadedFile('s.xlsx', output.getvalue())})
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'labour_strength')
		self.assertContains(response, 'quality_rating')
		self.assertFalse(DirectSubcontractor.objects.exists())

	def test_add_and_edit_in_portal(self):
		response = self.client.post('/subcontractors/add/', {
			'code': 'SC0100', 'name': 'Metro Composite Works', 'status': 'active', 'service_type': 'Composite',
			'labour_strength': 25, 'trades': 'Plumbing, Electrical', 'projects': 'Vignan Annapurna',
		})
		sub = DirectSubcontractor.objects.get(code='SC0100')
		self.assertRedirects(response, f'/subcontractors/{sub.pk}/')
		page = self.client.get(f'/subcontractors/{sub.pk}/edit/')
		self.assertContains(page, 'Labour Strength Available (Nos)')
		self.assertContains(page, 'value="Plumbing, Electrical"')
		labels = [field.label for field in page.context['form']][:7]
		self.assertEqual(labels, EXPECTED_COLUMNS[1:8])
		self.client.post(f'/subcontractors/{sub.pk}/edit/', {'code': 'SC0100', 'name': 'Metro Composite Works', 'status': 'on_hold', 'trades': 'Civil'})
		sub.refresh_from_db()
		self.assertEqual(sub.status, 'on_hold')
		self.assertEqual([link.trade.name for link in sub.trade_links.all()], ['Civil'])

	def test_read_only_user_cannot_change_data(self):
		viewer = get_user_model().objects.create_user('viewer', password='pw')
		self.client.force_login(viewer)
		self.assertEqual(self.client.get('/subcontractors/').status_code, 200)
		self.assertNotContains(self.client.get('/subcontractors/'), 'Upload subcontractors')
		for url in ('/subcontractors/add/', '/subcontractors/upload/', '/subcontractors/template.xlsx'):
			self.assertEqual(self.client.get(url).status_code, 302, url)

	def test_old_addresses_redirect(self):
		self.assertRedirects(self.client.get('/subcontractors/direct/'), '/subcontractors/')
		self.assertRedirects(self.client.get('/subcontractors/material-work/?q=x'), '/subcontractors/?q=x')


class SearchTests(TestCase):
	def setUp(self):
		self.client.force_login(get_user_model().objects.create_user('viewer', password='pw'))
		self.supplier = Supplier.objects.create(vendor_code='SP0009', name='Sri Lakshmi Traders', mobile_number='+91 98480 12345', city='Hyderabad')
		self.sub = DirectSubcontractor.objects.create(code='SC0009', name='Ganesh Labour Co', mobile_number='9848012345', alternate_number='040-2776 3835')
		DirectSubcontractor.objects.create(code='SC0010', name='Other Works', mobile_number='9000000000')

	def test_dashboard_search_finds_mobile_number_in_both_directories(self):
		for query in ('9848012345', '98480 12345', '+91-98480-12345'):
			response = self.client.get('/search/', {'q': query})
			self.assertContains(response, 'Sri Lakshmi Traders', msg_prefix=query)
			self.assertContains(response, 'Ganesh Labour Co', msg_prefix=query)
			self.assertNotContains(response, 'Other Works', msg_prefix=query)
			self.assertContains(response, 'Mobile No.', msg_prefix=query)

	def test_search_by_name_code_and_alternate_number(self):
		self.assertContains(self.client.get('/search/', {'q': 'SC0009'}), 'Ganesh Labour Co')
		self.assertContains(self.client.get('/search/', {'q': '27763835'}), 'Ganesh Labour Co')
		response = self.client.get('/search/', {'q': 'lakshmi'})
		self.assertContains(response, 'Sri Lakshmi Traders')
		self.assertNotContains(response, 'Ganesh Labour Co')

	def test_directory_lists_filter_by_mobile_number(self):
		self.assertContains(self.client.get('/suppliers/', {'q': '98480 12345'}), 'Sri Lakshmi Traders')
		response = self.client.get('/subcontractors/', {'q': '98480-12345'})
		self.assertContains(response, 'Ganesh Labour Co')
		self.assertNotContains(response, 'Other Works')

	def test_dashboard_has_search_box(self):
		response = self.client.get('/')
		self.assertContains(response, 'action="/search/"')
		self.assertContains(response, 'Labour &amp; Composite')
