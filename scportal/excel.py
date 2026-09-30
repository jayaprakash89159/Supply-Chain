"""Shared helpers for the Excel upload templates and imports."""
import re
from io import BytesIO
from zipfile import BadZipFile

from django.core.exceptions import ValidationError
from django.http import HttpResponse
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException

XLSX_CONTENT_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
READ_ERRORS = (BadZipFile, InvalidFileException, OSError, ValueError)


def header_key(value):
	return re.sub(r'[^a-z0-9]', '', str(value).casefold())


def parse_boolean(value):
	if value is None or value == '':
		return None
	if isinstance(value, bool):
		return value
	if isinstance(value, (int, float)) and value in (0, 1):
		return bool(value)
	key = str(value).strip().casefold()
	if key in {'yes', 'y', 'true', '1', 'ok', 'passed', 'valid'}:
		return True
	if key in {'no', 'n', 'false', '0', 'not ok', 'failed', 'invalid'}:
		return False
	raise ValidationError(f'Use Yes/No or OK/Not OK, not {value!r}.')


def parse_status(value):
	key = re.sub(r'[^a-z]', '', str(value).casefold())
	return {'active': 'active', 'onhold': 'on_hold', 'inactive': 'inactive'}.get(key, value)


def numbers_as_text(values, field_names):
	"""Excel stores phone, PIN and account numbers as numbers; keep them as text."""
	for field_name in field_names:
		value = values.get(field_name)
		if isinstance(value, (int, float)) and not isinstance(value, bool):
			values[field_name] = str(int(value))
	if str(values.get('gst_state_code', '')).isdigit():
		values['gst_state_code'] = str(values['gst_state_code']).zfill(2)


def template_response(filename, sheet_title, columns, sample_rows, non_text_columns, instructions):
	workbook = Workbook()
	sheet = workbook.active
	sheet.title = sheet_title
	sheet.append(columns)
	for row in sample_rows:
		sheet.append(row)
	sheet.freeze_panes = 'A2'
	sheet.auto_filter.ref = sheet.dimensions
	for cell in sheet[1]:
		cell.font = cell.font.copy(bold=True, color='FFFFFF')
		cell.fill = cell.fill.copy(fill_type='solid', fgColor='203B67')
		cell.alignment = cell.alignment.copy(wrap_text=True, vertical='center')
	sheet.row_dimensions[1].height = 42
	for column in sheet.columns:
		width = max(len(str(cell.value or '')) for cell in column)
		sheet.column_dimensions[column[0].column_letter].width = min(max(width + 2, 12), 32)
	for column_index, heading in enumerate(columns, start=1):
		if heading not in non_text_columns:
			for row_index in range(2, 1002):
				sheet.cell(row=row_index, column=column_index).number_format = '@'
	notes = workbook.create_sheet('Instructions')
	for line in instructions:
		notes.append([line])
	output = BytesIO()
	workbook.save(output)
	response = HttpResponse(output.getvalue(), content_type=XLSX_CONTENT_TYPE)
	response['Content-Disposition'] = f'attachment; filename="{filename}"'
	return response


def open_sheet(uploaded_file, preferred_names):
	"""Return (header_names, row_iterator) from the first matching sheet, else the first sheet."""
	try:
		workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
	except READ_ERRORS as exc:
		raise ValidationError(f'Could not read this Excel workbook: {exc}') from exc
	wanted = {header_key(name) for name in preferred_names}
	sheet_name = next((name for name in workbook.sheetnames if header_key(name) in wanted), workbook.sheetnames[0])
	rows = workbook[sheet_name].iter_rows(values_only=True)
	headers = next(rows, None)
	if not headers:
		raise ValidationError(f'The {sheet_name} sheet must have a header row.')
	header_names = [str(value).strip() if value is not None else '' for value in headers]
	nonempty = [name for name in header_names if name]
	if len(nonempty) != len(set(nonempty)):
		raise ValidationError('The header row contains duplicate column names.')
	return header_names, rows
