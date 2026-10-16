from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import Workbook, load_workbook

from materials.models import Brand, Material, MaterialCategory, Trade
from projects.models import Project
from relationships.models import SupplierBrand, SupplierMaterial, SupplierProject, SupplierTrade
from .models import Supplier


class SupplierDirectoryTests(TestCase):
	def setUp(self):
		self.user = get_user_model().objects.create_user(username='viewer', password='test-password')
		self.client.login(username='viewer', password='test-password')
		Supplier.objects.create(vendor_code='SUP-001', name='Northstar Steel', city='Pune', status='active')
		Supplier.objects.create(vendor_code='SUP-002', name='Riverbend Cement', city='Mumbai', status='on_hold')

	def test_directory_requires_authentication(self):
		self.client.logout()
		response = self.client.get('/suppliers/')
		self.assertRedirects(response, '/accounts/login/?next=/suppliers/')

	def test_search_filters_by_vendor_name(self):
		response = self.client.get('/suppliers/', {'q': 'northstar'})
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Northstar Steel')
		self.assertNotContains(response, 'Riverbend Cement')

	def test_status_filter_combines_with_search(self):
		response = self.client.get('/suppliers/', {'q': 'SUP-', 'status': 'active'})
		self.assertContains(response, 'Northstar Steel')
		self.assertNotContains(response, 'Riverbend Cement')

	def test_normal_user_cannot_open_admin_add_form(self):
		response = self.client.get('/admin/suppliers/supplier/add/')
		self.assertEqual(response.status_code, 302)
		self.assertTrue(response.url.startswith('/admin/login/'))

	def test_supplier_profile_renders(self):
		supplier = Supplier.objects.get(vendor_code='SUP-001')
		response = self.client.get(f'/suppliers/{supplier.pk}/')
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Northstar Steel')


class SupplierAdminImportTests(TestCase):
	def setUp(self):
		user = get_user_model().objects.create_superuser(
			username='admin', email='admin@example.com', password='test-password'
		)
		self.client.force_login(user)

	def workbook_upload(self, rows):
		workbook = Workbook()
		sheet = workbook.active
		sheet.title = 'Suppliers'
		sheet.append(['Vendor Code', 'Vendor / Firm Name', 'City', 'Overall Rating'])
		for row in rows:
			sheet.append(row)
		content = BytesIO()
		workbook.save(content)
		return SimpleUploadedFile(
			'suppliers.xlsx',
			content.getvalue(),
			content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
		)

	def test_admin_can_download_supplier_template(self):
		response = self.client.get('/admin/suppliers/supplier/template.xlsx')
		self.assertEqual(response.status_code, 200)
		workbook = load_workbook(BytesIO(response.content), read_only=True)
		self.assertIn('Instructions', workbook.sheetnames)
		rows = workbook['Suppliers'].iter_rows(values_only=True)
		headers = next(rows)
		sample = next(rows)
		self.assertIn('Vendor Code', headers)
		self.assertIn('Vendor / Firm Name', headers)
		self.assertEqual(len(sample), len(headers))
		self.assertEqual(sample[1:3], ('SP0001', 'Sri Lakshmi Traders'))
		response = self.client.post(
			'/admin/suppliers/supplier/import/',
			{'file': SimpleUploadedFile('supplier_template.xlsx', response.content)},
		)
		self.assertRedirects(response, '/admin/suppliers/supplier/')
		supplier = Supplier.objects.get(vendor_code='SP0001')
		self.assertEqual(supplier.gstin, '36BCNPP5410K1ZA')
		self.assertTrue(supplier.gstin_check)
		self.assertTrue(supplier.worked_with_vishwa_before)
		self.assertEqual(supplier.status, Supplier.Status.ACTIVE)
		self.assertIn('SRIKANTH PILLA', supplier.remarks)

	def test_admin_can_import_supplier_workbook(self):
		trade = Trade.objects.create(name='Electrical')
		category = MaterialCategory.objects.create(name='Electrical')
		material = Material.objects.create(name='Cable', category=category)
		brand = Brand.objects.create(name='Acme')
		project = Project.objects.create(code='P-001', name='Clinic Upgrade')
		workbook = Workbook()
		sheet = workbook.active
		sheet.title = 'Suppliers'
		sheet.append([
			'Vendor Code', 'Vendor / Firm Name', 'City', 'Overall Rating', 'Trade',
			'Material Category', 'Key Materials Supplied', 'Brands Dealt', 'Project(s) Worked On',
		])
		sheet.append([
			'SUP-003', 'Summit Hardware', 'Pune', 4.5, 'Electrical', 'Electrical',
			'Cable', 'Acme', 'P-001',
		])
		content = BytesIO()
		workbook.save(content)
		response = self.client.post(
			'/admin/suppliers/supplier/import/',
			{'file': SimpleUploadedFile('suppliers.xlsx', content.getvalue())},
		)
		self.assertRedirects(response, '/admin/suppliers/supplier/')
		supplier = Supplier.objects.get(vendor_code='SUP-003')
		self.assertEqual(supplier.name, 'Summit Hardware')
		self.assertEqual(supplier.city, 'Pune')
		self.assertEqual(supplier.overall_rating, 4.5)
		self.assertEqual(SupplierTrade.objects.get(supplier=supplier).trade, trade)
		self.assertEqual(SupplierMaterial.objects.get(supplier=supplier).material, material)
		self.assertEqual(SupplierBrand.objects.get(supplier=supplier).brand, brand)
		self.assertEqual(SupplierProject.objects.get(supplier=supplier).project, project)

	def test_invalid_workbook_data_is_rejected_without_importing_rows(self):
		response = self.client.post(
			'/admin/suppliers/supplier/import/',
			{'file': self.workbook_upload([
				['SUP-004', 'Valid Supplier', 'Pune', 4],
				['SUP-005', 'Invalid Supplier', 'Mumbai', 6],
			])},
		)
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'overall_rating')
		self.assertFalse(Supplier.objects.filter(vendor_code__in=['SUP-004', 'SUP-005']).exists())

	def test_import_creates_missing_master_data(self):
		workbook = Workbook()
		sheet = workbook.active
		sheet.title = 'Suppliers'
		sheet.append([
			'Vendor Code', 'Vendor / Firm Name', 'Trade', 'Material Category',
			'Key Materials Supplied', 'Supplier Type', 'Brands Dealt', 'Project(s) Worked On',
		])
		sheet.append(['SUP-010', 'New Traders', 'Civil, Plumbing', 'Cement', 'OPC 53, PPC', 'Dealer', 'UltraTech; ACC', 'Sreenidhi CR-2'])
		sheet.append(['SUP-011', 'Category Only', '', 'Steel', '', 'Distributor', '', ''])
		content = BytesIO()
		workbook.save(content)
		response = self.client.post('/admin/suppliers/supplier/import/', {'file': SimpleUploadedFile('s.xlsx', content.getvalue())})
		self.assertRedirects(response, '/admin/suppliers/supplier/')
		supplier = Supplier.objects.get(vendor_code='SUP-010')
		self.assertEqual(supplier.supplier_type, 'Dealer')
		self.assertEqual(sorted(link.trade.name for link in supplier.trade_links.all()), ['Civil', 'Plumbing'])
		self.assertEqual({(link.material.name, link.material.category.name) for link in supplier.material_links.all()}, {('OPC 53', 'Cement'), ('PPC', 'Cement')})
		self.assertEqual(sorted(link.brand.name for link in supplier.brand_links.all()), ['ACC', 'UltraTech'])
		self.assertEqual(supplier.project_links.get().project.name, 'Sreenidhi CR-2')
		other = Supplier.objects.get(vendor_code='SUP-011')
		self.assertEqual(other.material_category_links.get().category.name, 'Steel')
		profile = self.client.get(f'/suppliers/{supplier.pk}/')
		for text in ('Supply profile', 'Civil, Plumbing', 'Cement', 'OPC 53, PPC', 'Dealer', 'Sreenidhi CR-2'):
			self.assertContains(profile, text)

	def test_edit_form_shows_and_saves_linked_fields(self):
		supplier = Supplier.objects.create(vendor_code='SUP-020', name='Edit Me')
		response = self.client.post(f'/suppliers/{supplier.pk}/edit/', {
			'vendor_code': 'SUP-020', 'name': 'Edit Me', 'status': 'active', 'supplier_type': 'Manufacturer',
			'trades': 'Electrical', 'material_categories': 'Cables', 'materials': 'Armoured cable, Flexible wire',
			'brands': 'Polycab', 'projects': 'A S Rao Nagar',
		})
		self.assertRedirects(response, f'/suppliers/{supplier.pk}/')
		page = self.client.get(f'/suppliers/{supplier.pk}/edit/')
		self.assertContains(page, 'Key Materials Supplied')
		self.assertContains(page, 'Armoured cable, Flexible wire')
		self.assertContains(page, 'value="Polycab"')
		self.assertContains(page, 'value="A S Rao Nagar"')
		labels = [field.label for field in page.context['form']]
		self.assertEqual(labels[:8], ['Vendor Code', 'Vendor / Firm Name', 'Trade', 'Material Category', 'Key Materials Supplied', 'Supplier Type', 'Brands Dealt', 'Unit of Supply'])
		self.client.post(f'/suppliers/{supplier.pk}/edit/', {'vendor_code': 'SUP-020', 'name': 'Edit Me', 'status': 'active', 'trades': ''})
		self.assertFalse(supplier.trade_links.exists())


class GstinCheckTests(TestCase):
	def setUp(self):
		self.client.force_login(get_user_model().objects.create_superuser('gst', 'gst@example.com', 'pw'))

	def test_inspect_reads_state_pan_and_check_character(self):
		from scportal.gstin import inspect
		info = inspect(' 36bcnpp5410k1za ')
		self.assertEqual((info['status'], info['gstin'], info['state_code'], info['state'], info['pan']), ('ok', '36BCNPP5410K1ZA', '36', 'Telangana', 'BCNPP5410K'))
		typo = inspect('09AERPY3845G2Z0')
		self.assertEqual((typo['status'], typo['expected_check']), ('check_mismatch', 'O'))
		self.assertEqual(inspect('ABC')['status'], 'invalid_format')
		self.assertEqual(inspect('')['status'], 'empty')

	def test_form_tidies_gstin_and_rejects_bad_format(self):
		response = self.client.post('/suppliers/add/', {'vendor_code': 'SP0700', 'name': 'Tidy', 'status': 'active', 'gstin': '36bcnpp5410k1za', 'pan': 'bcnpp5410k'})
		supplier = Supplier.objects.get(vendor_code='SP0700')
		self.assertRedirects(response, f'/suppliers/{supplier.pk}/')
		self.assertEqual((supplier.gstin, supplier.pan), ('36BCNPP5410K1ZA', 'BCNPP5410K'))
		response = self.client.post('/suppliers/add/', {'vendor_code': 'SP0701', 'name': 'Bad', 'status': 'active', 'gstin': '36BCNPP5410'})
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Not a valid GSTIN format')
		self.assertFalse(Supplier.objects.filter(vendor_code='SP0701').exists())

	def test_check_mismatch_is_saved_but_flagged_on_profile(self):
		self.client.post('/suppliers/add/', {'vendor_code': 'SP0702', 'name': 'Typo', 'status': 'active', 'gstin': '09AERPY3845G2Z0'})
		supplier = Supplier.objects.get(vendor_code='SP0702')
		page = self.client.get(f'/suppliers/{supplier.pk}/')
		self.assertContains(page, 'Check character mismatch')
		self.assertContains(page, 'data-gst-check="09AERPY3845G2Z0"')
		self.assertContains(page, 'https://services.gst.gov.in/services/searchtp')
		self.assertContains(page, 'js/gstin.js')

	def test_form_pages_load_gstin_helper(self):
		for url in ('/suppliers/add/', '/subcontractors/add/'):
			page = self.client.get(url)
			self.assertContains(page, 'id="gst-config"', msg_prefix=url)
			self.assertContains(page, '"36": "Telangana"', msg_prefix=url)
