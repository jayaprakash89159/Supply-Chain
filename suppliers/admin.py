import re
from io import BytesIO
from zipfile import BadZipFile

from django import forms
from django.contrib import admin, messages
from django.core.exceptions import ValidationError
from django.db import transaction
from django.http import HttpResponse
from django.shortcuts import redirect, render
from django.urls import path, reverse
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .links import LINK_FIELDS, save_links, split_names
from .models import Supplier


UPLOAD_COLUMNS = [
	'Sl No', 'Vendor Code', 'Vendor / Firm Name', 'Trade', 'Material Category',
	'Key Materials Supplied', 'Supplier Type', 'Brands Dealt', 'Unit of Supply',
	'Contact Person', 'Mobile No.', 'Alternate No.', 'Email', 'Full Address', 'City',
	'State', 'PIN Code', 'GST Registration Type', 'GSTIN', 'GST State Code', 'PAN',
	'GSTIN Check', 'MSME Category', 'Udyam Reg. No.', 'TDS Section',
	'Bank A/c Holder Name', 'Bank Name', 'Bank A/c No.', 'IFSC', 'Branch',
	'IFSC Check', 'Referred By', 'Worked with VISHWA Before?', 'Project(s) Worked On',
	'Quality (1-5)', 'Delivery (1-5)', 'Price (1-5)', 'Overall Rating',
	'Last Reviewed On', 'Status', 'Remarks',
]

FIELD_HEADERS = {
	'Vendor Code': 'vendor_code',
	'Vendor / Firm Name': 'name',
	'Supplier Type': 'supplier_type',
	'Unit of Supply': 'unit_of_supply',
	'Contact Person': 'contact_person',
	'Mobile No.': 'mobile_number',
	'Alternate No.': 'alternate_number',
	'Email': 'email',
	'Full Address': 'full_address',
	'City': 'city',
	'State': 'state',
	'PIN Code': 'pin_code',
	'GST Registration Type': 'gst_registration_type',
	'GSTIN': 'gstin',
	'GST State Code': 'gst_state_code',
	'PAN': 'pan',
	'GSTIN Check': 'gstin_check',
	'MSME Category': 'msme_category',
	'Udyam Reg. No.': 'udyam_registration_number',
	'TDS Section': 'tds_section',
	'Bank A/c Holder Name': 'bank_account_holder_name',
	'Bank Name': 'bank_name',
	'Bank A/c No.': 'bank_account_number',
	'IFSC': 'ifsc',
	'Branch': 'bank_branch',
	'IFSC Check': 'ifsc_check',
	'Referred By': 'referred_by',
	'Worked with VISHWA Before?': 'worked_with_vishwa_before',
	'Quality (1-5)': 'quality_rating',
	'Delivery (1-5)': 'delivery_rating',
	'Price (1-5)': 'price_rating',
	'Overall Rating': 'overall_rating',
	'Last Reviewed On': 'last_reviewed_on',
	'Status': 'status',
	'Remarks': 'remarks',
}


def _header_key(value):
	return re.sub(r'[^a-z0-9]', '', str(value).casefold())


class SupplierUploadForm(forms.Form):
	file = forms.FileField(label='Supplier workbook (.xlsx)')


@admin.register(Supplier)
class SupplierAdmin(admin.ModelAdmin):
	list_display = ('vendor_code', 'name', 'supplier_type', 'city', 'state', 'status', 'overall_rating')
	list_filter = ('status', 'supplier_type', 'city', 'state', 'msme_category')
	search_fields = ('vendor_code', 'name', 'gstin', 'pan', 'contact_person', 'city', 'state')
	readonly_fields = ('created_at', 'updated_at')
	list_per_page = 50
	change_list_template = 'admin/suppliers/supplier/change_list.html'

	def get_urls(self):
		custom_urls = [
			path('import/', self.admin_site.admin_view(self.import_view), name='suppliers_supplier_import'),
			path('template.xlsx', self.admin_site.admin_view(self.template_download), name='suppliers_supplier_template'),
		]
		return custom_urls + super().get_urls()

	def template_download(self, request):
		workbook = Workbook()
		sheet = workbook.active
		sheet.title = 'Suppliers'
		sheet.append(UPLOAD_COLUMNS)
		sample = [
			1, 'SP0001', 'Sri Lakshmi Traders', '', '', '', '', '', '', '', '', '', '',
			'4-4-214/A-5Vimalplaza', 'Ranigunj', 'Telangana', '500003', 'Regular',
			'36BCNPP5410K1ZA', '36', 'BCNPP5410K', 'OK', '', '', '', '', '', '', '',
			'', '', '', 'Yes', '', '', '', '', '', '', 'Active',
			'Legal name: SRIKANTH PILLA; ERP',
		]
		sheet.append(sample)
		sheet.append([
			2, 'SP0002', 'Shiram Shakti Enterprises', '', '', '', '', '', '', '', '', '', '',
			'Gandhi Gram,Kailsh Kunj,Lucknow', 'Lucknow', 'Uttar Pradesh', '226016', 'Regular',
			'09AERPY3845G2Z0', '09', 'AERPY3845G', 'OK', '', '', '', '', '', '', '',
			'', '', '', 'Yes', '', '', '', '', '', '', 'Active',
			'Legal name: AMARNATH YADAV; ERP',
		])
		sheet.freeze_panes = 'A2'
		sheet.auto_filter.ref = sheet.dimensions
		for cell in sheet[1]:
			cell.font = cell.font.copy(bold=True, color='FFFFFF')
			cell.fill = cell.fill.copy(fill_type='solid', fgColor='203B67')
			cell.alignment = cell.alignment.copy(wrap_text=True, vertical='center')
		sheet.row_dimensions[1].height = 42
		for column in sheet.columns:
			column_letter = column[0].column_letter
			max_length = max(len(str(cell.value or '')) for cell in column)
			sheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 32)
		for column_index, heading in enumerate(UPLOAD_COLUMNS, start=1):
			if heading not in {'Sl No', 'GSTIN Check', 'IFSC Check', 'Worked with VISHWA Before?', 'Quality (1-5)', 'Delivery (1-5)', 'Price (1-5)', 'Overall Rating', 'Last Reviewed On'}:
				for row_index in range(2, 1002):
					sheet.cell(row=row_index, column=column_index).number_format = '@'
		instructions = workbook.create_sheet('Instructions')
		instructions.append(['Supplier upload template'])
		instructions.append(['The Suppliers sheet uses the supplied spreadsheet columns and includes two example rows (SP0001, SP0002). Replace or delete them before a real upload.'])
		instructions.append(['Required columns: Vendor Code and Vendor / Firm Name.'])
		instructions.append(['Trade, Material Category, Key Materials Supplied, Brands Dealt and Project(s) Worked On accept any names. Separate multiple names with commas or semicolons. New names are added to the portal master data automatically.'])
		instructions.append(['Give one Material Category for all materials, or one category per material in the same order. Materials without a category are filed under General.'])
		instructions.append(['Status: Active, On hold, or Inactive. Checks accept OK/Not OK or Yes/No. Ratings must be between 0 and 5.'])
		instructions.append(['Dates use YYYY-MM-DD. Existing vendor codes cannot be uploaded again.'])
		output = BytesIO()
		workbook.save(output)
		response = HttpResponse(
			output.getvalue(),
			content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
		)
		response['Content-Disposition'] = 'attachment; filename="supplier_upload_template.xlsx"'
		return response

	def import_view(self, request):
		form = SupplierUploadForm(request.POST or None, request.FILES or None)
		if request.method == 'POST' and form.is_valid():
			try:
				created_count = self._import_workbook(form.cleaned_data['file'])
			except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
				form.add_error('file', f'Could not read this Excel workbook: {exc}')
			except ValidationError as exc:
				form.add_error('file', exc.messages)
			else:
				self.message_user(request, f'Imported {created_count} supplier(s).', messages.SUCCESS)
				return redirect(reverse('admin:suppliers_supplier_changelist'))
		context = {
			**self.admin_site.each_context(request),
			'opts': self.model._meta,
			'form': form,
			'title': 'Upload suppliers from Excel',
		}
		return render(request, 'admin/suppliers/supplier/import.html', context)

	def _import_workbook(self, uploaded_file):
		try:
			workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
		except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
			raise ValidationError(f'Could not read this Excel workbook: {exc}') from exc
		if 'Suppliers' not in workbook.sheetnames:
			raise ValidationError('The workbook must contain a sheet named Suppliers.')

		sheet = workbook['Suppliers']
		rows = sheet.iter_rows(values_only=True)
		headers = next(rows, None)
		if not headers:
			raise ValidationError('The Suppliers sheet must have a header row.')
		header_names = [str(value).strip() if value is not None else '' for value in headers]
		nonempty_headers = [name for name in header_names if name]
		if len(nonempty_headers) != len(set(nonempty_headers)):
			raise ValidationError('The header row contains duplicate column names.')

		field_map = {_header_key(header): field for header, field in FIELD_HEADERS.items()}
		for field in Supplier._meta.fields:
			if field.editable and field.name not in {'created_at', 'updated_at', 'archived_at'}:
				field_map.setdefault(_header_key(field.name), field.name)
		field_map[_header_key('Sl No')] = None
		field_map[_header_key('Trade')] = 'trades'
		field_map[_header_key('Material Category')] = 'material_categories'
		field_map[_header_key('Key Materials Supplied')] = 'materials'
		field_map[_header_key('Brands Dealt')] = 'brands'
		field_map[_header_key('Project(s) Worked On')] = 'projects'
		unknown_headers = [name for name in header_names if name and _header_key(name) not in field_map]
		if unknown_headers:
			raise ValidationError(f'Unknown column(s): {", ".join(unknown_headers)}.')
		mapped_headers = {_header_key(name) for name in header_names}
		missing_required = {
			name for name in ('vendor_code', 'name')
			if not any(field == name and key in mapped_headers for key, field in field_map.items())
		}
		if missing_required:
			raise ValidationError(f'Missing required column(s): {", ".join(sorted(missing_required))}.')

		import_rows = []
		seen_vendor_codes = set()
		row_errors = []
		for row_number, row in enumerate(rows, start=2):
			if not any(value is not None and str(value).strip() for value in row):
				continue
			values = {}
			link_names = {field: [] for field in LINK_FIELDS}
			for header, value in zip(header_names, row):
				if not header or value is None:
					continue
				mapped_field = field_map[_header_key(header)]
				if mapped_field is None:
					continue
				if mapped_field in link_names:
					link_names[mapped_field] = split_names(value)
				else:
					values[mapped_field] = value.strip() if isinstance(value, str) else value
			row_error_count = len(row_errors)
			for field_name in ('gstin_check', 'ifsc_check', 'worked_with_vishwa_before'):
				if field_name in values:
					try:
						values[field_name] = self._parse_boolean(values[field_name])
					except ValidationError as exc:
						row_errors.extend(f'Row {row_number}, {field_name}: {error}' for error in exc.messages)
			for field_name in ('pin_code', 'mobile_number', 'alternate_number', 'bank_account_number', 'gst_state_code'):
				if isinstance(values.get(field_name), (int, float)) and not isinstance(values[field_name], bool):
					values[field_name] = str(int(values[field_name]))
			if str(values.get('gst_state_code', '')).isdigit():
				values['gst_state_code'] = values['gst_state_code'].zfill(2)
			if 'status' in values:
				values['status'] = self._parse_status(values['status'])
			vendor_code = str(values.get('vendor_code', '')).strip()
			if vendor_code in seen_vendor_codes:
				row_errors.append(f'Row {row_number}: duplicate vendor_code {vendor_code!r} in this workbook.')
				continue
			seen_vendor_codes.add(vendor_code)
			supplier = Supplier(**values)
			try:
				supplier.full_clean()
			except ValidationError as exc:
				messages_by_field = getattr(exc, 'message_dict', {'row': exc.messages})
				row_errors.extend(
					f'Row {row_number}, {field}: {", ".join(errors)}'
					for field, errors in messages_by_field.items()
				)
			if len(row_errors) == row_error_count:
				import_rows.append((supplier, link_names))

		if row_errors:
			raise ValidationError(row_errors)
		if not import_rows:
			raise ValidationError('The Suppliers sheet contains no supplier rows to import.')
		with transaction.atomic():
			for supplier, link_names in import_rows:
				supplier.save()
				save_links(supplier, link_names)
		return len(import_rows)

	@staticmethod
	def _parse_boolean(value):
		if value is None or value == '':
			return None
		if isinstance(value, bool):
			return value
		if isinstance(value, (int, float)) and value in (0, 1):
			return bool(value)
		key = str(value).strip().casefold()
		if key in {'yes', 'y', 'true', '1', 'ok', 'passed', 'valid'}:
			return True
		if key in {'no', 'n', 'false', '0', 'not ok', 'failed', 'invalid'}:
			return False
		raise ValidationError(f'Use Yes/No or OK/Not OK, not {value!r}.')

	@staticmethod
	def _parse_status(value):
		key = re.sub(r'[^a-z]', '', str(value).casefold())
		return {'active': 'active', 'onhold': 'on_hold', 'inactive': 'inactive'}.get(key, value)
